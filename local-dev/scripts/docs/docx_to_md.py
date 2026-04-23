#!/usr/bin/env python3
"""Convert a .docx file to Markdown.

Handles paragraphs, headings, lists, tables, bold/italic,
and auto-detects HL7 messages and JSON blocks to wrap them in code fences.

Usage:
    python docx_to_md.py --input doc.docx [--output doc.md]

If --output is omitted, prints to stdout.
"""

import sys as _sys, os as _os; _sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), '..', 'lib'))
import argparse
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


# HL7 segment names that start a line
HL7_SEGMENTS = re.compile(
    r'^(MSH|PID|PV1|PV2|EVN|NK1|AL1|GT1|IN1|IN2|ORC|OBR|OBX|RXA|RXE|RXD|RXO|RXR|TQ1|DG1|PR1|FT1|ZPD|MRG|SCH|AIG|AIL|AIP|AIS|ADT|SFT|UAC)\|'
)

# Style-name to markdown heading level
HEADING_MAP = {
    'Title': '#',
    'Heading 1': '#',
    'Heading 2': '##',
    'Heading 3': '###',
    'Heading 4': '####',
    'Heading 5': '#####',
    'Heading 6': '######',
}


def runs_to_md(paragraph):
    """Convert paragraph runs to markdown with bold/italic."""
    parts = []
    for run in paragraph.runs:
        text = run.text
        if not text:
            continue
        if run.bold and run.italic:
            text = f'***{text}***'
        elif run.bold:
            text = f'**{text}**'
        elif run.italic:
            text = f'*{text}*'
        parts.append(text)
    # Fall back to plain text if runs are empty but paragraph has text
    if not parts and paragraph.text:
        return paragraph.text
    return ''.join(parts)


def table_to_md(table):
    """Convert a docx table to a markdown table."""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace('|', '\\|') for cell in row.cells]
        rows.append(cells)

    if not rows:
        return ''

    lines = []
    # Header row
    lines.append('| ' + ' | '.join(rows[0]) + ' |')
    # Separator
    lines.append('| ' + ' | '.join(['---'] * len(rows[0])) + ' |')
    # Data rows
    for row in rows[1:]:
        # Pad row if needed
        while len(row) < len(rows[0]):
            row.append('')
        lines.append('| ' + ' | '.join(row[:len(rows[0])]) + ' |')

    return '\n'.join(lines)


def is_json_line(text):
    """Check if a line looks like part of a JSON block."""
    stripped = text.strip()
    return stripped in ('{', '}', '},', '],', '[', ']') or \
        stripped.startswith('"') or \
        stripped.startswith('{') or \
        stripped.startswith('}') or \
        stripped.startswith('[') or \
        stripped.startswith(']')


def convert(input_path, output_path=None):
    """Convert a docx file to markdown."""
    doc = Document(input_path)

    # Build an ordered list of block elements (paragraphs and tables)
    # by walking the document body XML to preserve order
    from docx.oxml.ns import qn
    blocks = []
    table_index = 0
    para_index = 0
    tables = doc.tables
    paragraphs = doc.paragraphs

    for element in doc.element.body:
        if element.tag == qn('w:p'):
            if para_index < len(paragraphs):
                blocks.append(('para', paragraphs[para_index]))
                para_index += 1
        elif element.tag == qn('w:tbl'):
            if table_index < len(tables):
                blocks.append(('table', tables[table_index]))
                table_index += 1

    lines = []
    in_hl7_block = False
    in_json_block = False
    json_brace_depth = 0

    for block_type, block in blocks:
        if block_type == 'table':
            # Close any open code blocks
            if in_hl7_block:
                lines.append('```')
                in_hl7_block = False
            if in_json_block:
                lines.append('```')
                in_json_block = False
                json_brace_depth = 0

            lines.append('')
            lines.append(table_to_md(block))
            lines.append('')
            continue

        # It's a paragraph
        para = block
        text = para.text.strip()

        if not text:
            # Close HL7 block on empty line
            if in_hl7_block:
                lines.append('```')
                in_hl7_block = False
            # Close JSON block on empty line only if braces are balanced
            if in_json_block and json_brace_depth <= 0:
                lines.append('```')
                in_json_block = False
                json_brace_depth = 0
            continue

        style_name = para.style.name if para.style else 'Normal'

        # Check for HL7 segment lines
        if HL7_SEGMENTS.match(text):
            if not in_hl7_block:
                if in_json_block:
                    lines.append('```')
                    in_json_block = False
                    json_brace_depth = 0
                lines.append('')
                lines.append('```hl7')
                in_hl7_block = True
            lines.append(text)
            continue

        # Close HL7 block if we hit a non-HL7 line
        if in_hl7_block:
            lines.append('```')
            in_hl7_block = False

        # Check for JSON blocks
        if is_json_line(text) and not in_hl7_block:
            if not in_json_block:
                lines.append('')
                lines.append('```json')
                in_json_block = True
                json_brace_depth = 0
            json_brace_depth += text.count('{') + text.count('[')
            json_brace_depth -= text.count('}') + text.count(']')
            lines.append(text)
            # Close JSON block if braces are balanced
            if json_brace_depth <= 0 and text.strip() == '}':
                lines.append('```')
                in_json_block = False
                json_brace_depth = 0
            continue

        # Close JSON block if we hit a non-JSON line
        if in_json_block:
            lines.append('```')
            in_json_block = False
            json_brace_depth = 0

        # Headings
        if style_name in HEADING_MAP:
            prefix = HEADING_MAP[style_name]
            lines.append('')
            lines.append(f'{prefix} {text}')
            lines.append('')
            continue

        # List paragraphs
        if style_name == 'List Paragraph':
            lines.append(f'- {runs_to_md(para)}')
            continue

        # Regular paragraphs
        md_text = runs_to_md(para)
        lines.append('')
        lines.append(md_text)

    # Close any trailing open blocks
    if in_hl7_block:
        lines.append('```')
    if in_json_block:
        lines.append('```')

    # Clean up: collapse multiple blank lines
    result = []
    prev_blank = False
    for line in lines:
        if line == '':
            if not prev_blank:
                result.append('')
            prev_blank = True
        else:
            prev_blank = False
            result.append(line)

    output = '\n'.join(result).strip() + '\n'

    if output_path:
        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(output)
        print(f'Converted: {input_path} -> {output_path}', file=sys.stderr)
        print(f'Size: {len(output)} bytes, {output.count(chr(10))} lines', file=sys.stderr)
    else:
        print(output)


def main():
    parser = argparse.ArgumentParser(description='Convert .docx to Markdown')
    parser.add_argument('--input', required=True, help='Path to .docx file')
    parser.add_argument('--output', help='Output .md file path (default: stdout)')
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f'Error: File not found: {args.input}', file=sys.stderr)
        sys.exit(1)

    if not args.input.lower().endswith('.docx'):
        print(f'Error: File must be a .docx file: {args.input}', file=sys.stderr)
        sys.exit(1)

    convert(args.input, args.output)


if __name__ == '__main__':
    main()
